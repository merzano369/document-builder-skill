"""
DocumentStyles.py
Рушій Word OpenXML для генерації документів.
Конфігурація стилів повністю відокремлена і завантажується з файлу styles_config.json.
"""

from __future__ import annotations
import json
import os
import re
import tempfile
import uuid
from typing import List, Optional, Dict, Any, Tuple

# Опціональні модулі для нативної конвертації LaTeX / MathML в OpenXML oMath
try:
    import lxml.etree as etree
    HAS_LXML = True
except ImportError:
    etree = None
    HAS_LXML = False

try:
    import latex2mathml.converter
    HAS_LATEX2MATHML = True
except ImportError:
    latex2mathml = None
    HAS_LATEX2MATHML = False

try:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    HAS_MATPLOTLIB = True
except ImportError:
    plt = None
    HAS_MATPLOTLIB = False

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

# Конфігурація за замовчуванням (використовується, якщо файл styles_config.json відсутній)
DEFAULT_CONFIG: Dict[str, Any] = {
    "font_name": "Times New Roman",
    "font_size_body": 14.0,
    "font_size_h1": 16.0,
    "font_size_h2": 14.0,
    "font_size_h3": 14.0,
    "paragraph_align": "justify",
    "bullet_align": "justify",
    "numbered_align": "justify",
    "heading_align": "left",
    "heading_h1_align": "left",
    "heading_h2_align": "left",
    "heading_h3_align": "left",
    "h1_page_break_before": True,
    "table_caption_position": "above",
    "table_caption_align": "center",
    "image_caption_position": "below",
    "image_caption_align": "center",
    "line_spacing": 1.15,
    "paragraph_indent_cm": 1.25,
    "margin_left_cm": 3.0,
    "margin_right_cm": 1.0,
    "margin_top_cm": 2.0,
    "margin_bottom_cm": 2.0,
    "page_numbering": True,
    "page_number_align": "center",
    "page_number_size": 12.0,
    "page_number_diff_first": True,
    "table_border_color": "000000",
    "table_border_sz": "8",
    "images_dir": "images",
}


def clean_hex(hex_str: str) -> str:
    """Очищує hex-колір від символу '#'."""
    return hex_str.lstrip("#").upper()


def hex_to_rgb(hex_str: str) -> RGBColor:
    """Конвертує Hex рядок в docx RGBColor."""
    c = clean_hex(hex_str)
    if len(c) != 6:
        c = "000000"
    return RGBColor(int(c[0:2], 16), int(c[2:4], 16), int(c[4:6], 16))


# Регулярний вираз для базового вбудованого форматування Markdown та формул:
# `код`, $inline_math$, ***жирний курсив***, **жирний**, *курсив*
MD_INLINE_RE = re.compile(
    r'(?P<code>`[^`\n]+`)'
    r'|(?P<math>\$(?!\s)[^$\n]+?(?<!\s)\$)'
    r'|(?P<bold_italic>\*\*\*(?!\s)[^*\n]+?(?<!\s)\*\*\*)'
    r'|(?P<bold>\*\*(?!\s)[^*\n]+?(?<!\s)\*\*)'
    r'|(?P<italic>\*(?!\s)[^*\n]+?(?<!\s)\*)'
)

ALIGN_MAP = {
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
}


class DocumentStyles:
    """
    Головний клас-рушій оформлення документа на базі OpenXML.
    Автоматично зчитує налаштування з styles_config.json та надає чистий API для коду.
    """

    def __init__(
        self,
        config_path: Optional[str] = None,
        doc_path: Optional[str] = None,
    ):
        self.config = self._load_config(config_path)

        # Зручні посилання на основні параметри
        self.font_name: str = self.config.get("font_name", "Times New Roman")
        self.font_size_body: float = float(self.config.get("font_size_body", 14.0))
        self.font_size_h1: float = float(self.config.get("font_size_h1", 16.0))
        self.font_size_h2: float = float(self.config.get("font_size_h2", 14.0))
        self.font_size_h3: float = float(self.config.get("font_size_h3", 14.0))

        # Вирівнювання тексту, списків та заголовків
        self.paragraph_align: str = str(self.config.get("paragraph_align", "justify"))
        self.bullet_align: str = str(self.config.get("bullet_align", "justify"))
        self.numbered_align: str = str(self.config.get("numbered_align", "justify"))
        self.heading_align: str = str(self.config.get("heading_align", "left"))
        self.heading_h1_align: str = str(self.config.get("heading_h1_align", self.heading_align))
        self.heading_h2_align: str = str(self.config.get("heading_h2_align", self.heading_align))
        self.heading_h3_align: str = str(self.config.get("heading_h3_align", self.heading_align))

        # Сторінкові розриви перед H1
        self.h1_page_break_before: bool = bool(self.config.get("h1_page_break_before", True))
        self._h1_count: int = 0

        # Розташування та вирівнювання підписів таблиць та фото
        self.table_caption_position: str = str(self.config.get("table_caption_position", "above")).lower()
        self.table_caption_align: str = str(self.config.get("table_caption_align", "center")).lower()
        self.image_caption_position: str = str(self.config.get("image_caption_position", "below")).lower()
        self.image_caption_align: str = str(self.config.get("image_caption_align", "center")).lower()

        self.line_spacing: float = float(self.config.get("line_spacing", 1.15))
        self.paragraph_indent_cm: float = float(self.config.get("paragraph_indent_cm", 1.25))

        self.margin_left_cm: float = float(self.config.get("margin_left_cm", 3.0))
        self.margin_right_cm: float = float(self.config.get("margin_right_cm", 1.0))
        self.margin_top_cm: float = float(self.config.get("margin_top_cm", 2.0))
        self.margin_bottom_cm: float = float(self.config.get("margin_bottom_cm", 2.0))

        self.table_border_color: str = str(self.config.get("table_border_color", "000000"))
        self.table_border_sz: str = str(self.config.get("table_border_sz", "8"))
        self.images_dir: str = str(self.config.get("images_dir", "images"))

        # Ініціалізація документа
        self.doc = Document(doc_path) if doc_path else Document()
        self._setup_page_margins()
        self._setup_default_styles()

        if self.config.get("page_numbering", True):
            self._setup_page_numbering()

    def _load_config(self, config_path: Optional[str]) -> Dict[str, Any]:
        """
        Завантажує конфігурацію з JSON файлу або повертає дефолтну.
        """
        cfg = dict(DEFAULT_CONFIG)

        target_path = config_path
        if not target_path:
            # Шукаємо styles_config.json у папці виконання або у папці скрипта
            candidates = [
                "styles_config.json",
                os.path.join(os.path.dirname(os.path.abspath(__file__)), "styles_config.json"),
            ]
            for c in candidates:
                if os.path.exists(c):
                    target_path = c
                    break

        if target_path and os.path.exists(target_path):
            try:
                with open(target_path, "r", encoding="utf-8") as f:
                    loaded = json.load(f)
                    if isinstance(loaded, dict):
                        for k, v in loaded.items():
                            if isinstance(v, dict):
                                cfg.update(v)
                            else:
                                cfg[k] = v
            except Exception as e:
                print(f"⚠️ Попередження: Не вдалося прочитати {target_path} ({e}), використано дефолтні стилі.")

        return cfg

    def _setup_page_margins(self):
        """Встановлює розміри полів документа."""
        for section in self.doc.sections:
            section.top_margin = Cm(self.margin_top_cm)
            section.bottom_margin = Cm(self.margin_bottom_cm)
            section.left_margin = Cm(self.margin_left_cm)
            section.right_margin = Cm(self.margin_right_cm)

    def _setup_default_styles(self):
        """Налаштовує дефолтні стилі Normal, List Bullet та List Number у Word."""
        normal_style = self.doc.styles["Normal"]
        normal_style.font.name = self.font_name
        normal_style.font.size = Pt(self.font_size_body)
        normal_style.font.color.rgb = hex_to_rgb("000000")
        normal_style.paragraph_format.line_spacing = self.line_spacing
        normal_style.paragraph_format.space_before = Pt(0)
        normal_style.paragraph_format.space_after = Pt(0)
        normal_style.paragraph_format.alignment = ALIGN_MAP.get(self.paragraph_align.lower(), WD_ALIGN_PARAGRAPH.JUSTIFY)

        if "List Bullet" in self.doc.styles:
            lb_style = self.doc.styles["List Bullet"]
            lb_style.font.name = self.font_name
            lb_style.font.size = Pt(self.font_size_body)
            lb_style.font.color.rgb = hex_to_rgb("000000")
            lb_style.paragraph_format.alignment = ALIGN_MAP.get(self.bullet_align.lower(), WD_ALIGN_PARAGRAPH.JUSTIFY)

        if "List Number" in self.doc.styles:
            ln_style = self.doc.styles["List Number"]
            ln_style.font.name = self.font_name
            ln_style.font.size = Pt(self.font_size_body)
            ln_style.font.color.rgb = hex_to_rgb("000000")
            ln_style.paragraph_format.alignment = ALIGN_MAP.get(self.numbered_align.lower(), WD_ALIGN_PARAGRAPH.JUSTIFY)

        # Налаштування нативних стилів заголовків Word (Heading 1 - Heading 3)
        heading_sizes = {1: self.font_size_h1, 2: self.font_size_h2, 3: self.font_size_h3}
        for lvl in range(1, 4):
            h_name = f"Heading {lvl}"
            if h_name in self.doc.styles:
                h_style = self.doc.styles[h_name]
                h_style.font.name = self.font_name
                h_style.font.size = Pt(heading_sizes.get(lvl, self.font_size_h2))
                h_style.font.bold = True
                h_style.font.color.rgb = hex_to_rgb("000000")
                h_style.paragraph_format.line_spacing = self.line_spacing
                h_style.paragraph_format.keep_with_next = True

    def _setup_page_numbering(self):
        """Вмикає нативну динамічну нумерацію сторінок у нижньому колонтитулі."""
        section = self.doc.sections[0]
        if self.config.get("page_number_diff_first", True):
            section.different_first_page_header_footer = True

        footer = section.footer
        p = footer.paragraphs[0]
        p.text = ""

        align = self.config.get("page_number_align", "center")
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if align == "right" else WD_ALIGN_PARAGRAPH.CENTER

        r_pg = p.add_run()
        num_size = float(self.config.get("page_number_size", 12.0))
        self._apply_run_font(r_pg, font_name=self.font_name, size_pt=num_size, color_hex="000000")
        self._insert_field_code(r_pg, "PAGE")

    def _apply_run_font(
        self,
        run,
        font_name: Optional[str] = None,
        size_pt: Optional[float] = None,
        bold: bool = False,
        italic: bool = False,
        color_hex: str = "000000",
    ):
        font = run.font
        font.name = font_name or self.font_name
        font.size = Pt(size_pt or self.font_size_body)
        font.bold = bold
        font.italic = italic
        font.color.rgb = hex_to_rgb(color_hex)

    def _insert_field_code(self, run, field_code: str):
        """Вставляє динамічне OpenXML поле (PAGE, NUMPAGES, TOC тощо)."""
        fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls("w"))
        instrText = parse_xml(
            r'<w:instrText %s xml:space="preserve"> %s </w:instrText>'
            % (nsdecls("w"), field_code.strip())
        )
        fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls("w"))
        fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls("w"))
        run._r.append(fldChar1)
        run._r.append(instrText)
        run._r.append(fldChar2)
        run._r.append(fldChar3)

    def _init_math_engine(self):
        """Ініціалізує XSLT перетворювач MathML -> OMML (MML2OMML.XSL)."""
        if hasattr(self, "_math_engine_initialized") and self._math_engine_initialized:
            return
        self._math_engine_initialized = True
        self._xslt = None

        if not HAS_LXML:
            return

        candidates = [
            os.path.join(os.path.dirname(os.path.abspath(__file__)), "MML2OMML.XSL"),
            "MML2OMML.XSL",
            r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL",
            r"C:\Program Files (x86)\Microsoft Office\root\Office16\MML2OMML.XSL",
            r"C:\Program Files\Microsoft Office\Office16\MML2OMML.XSL",
            r"C:\Program Files\Microsoft Office\Office15\MML2OMML.XSL",
            r"C:\Program Files\Microsoft Office\Office14\MML2OMML.XSL",
        ]
        for c in candidates:
            if os.path.exists(c):
                try:
                    tree = etree.parse(c)
                    self._xslt = etree.XSLT(tree)
                    break
                except Exception:
                    pass

    def convert_to_omml(self, formula: str) -> Optional[str]:
        """
        Конвертує LaTeX або MathML рядок у валідний XML рядок <m:oMath>.
        Повертає None у разі неможливості конвертації.
        """
        self._init_math_engine()
        if not self._xslt:
            return None

        clean = formula.strip()
        if clean.startswith("$$") and clean.endswith("$$"):
            clean = clean[2:-2].strip()
        elif clean.startswith("$") and clean.endswith("$"):
            clean = clean[1:-1].strip()

        try:
            if clean.startswith("<math") or "<math" in clean:
                mml_str = clean
            else:
                if not HAS_LATEX2MATHML:
                    return None
                mml_str = latex2mathml.converter.convert(clean)

            mml_tree = etree.fromstring(mml_str)
            omml_tree = self._xslt(mml_tree)
            return etree.tostring(omml_tree, encoding="utf-8").decode("utf-8")
        except Exception:
            return None

    def _render_latex_to_temp_png(self, formula: str) -> Optional[str]:
        """Рендерить LaTeX формулу у тимчасовий прозорий PNG високої чіткості (fallback)."""
        if not HAS_MATPLOTLIB:
            return None
        clean = formula.strip()
        if clean.startswith("$$") and clean.endswith("$$"):
            clean = clean[2:-2].strip()
        elif clean.startswith("$") and clean.endswith("$"):
            clean = clean[1:-1].strip()

        try:
            temp_dir = tempfile.gettempdir()
            filename = f"formula_{uuid.uuid4().hex[:8]}.png"
            temp_path = os.path.join(temp_dir, filename)

            fig = plt.figure()
            fig.text(0.5, 0.5, f"${clean}$", fontsize=16, ha="center", va="center")
            fig.patch.set_visible(False)
            plt.axis("off")
            fig.savefig(temp_path, bbox_inches="tight", pad_inches=0.05, dpi=300, transparent=True)
            plt.close(fig)
            return temp_path
        except Exception:
            return None

    def _add_formatted_text(
        self,
        paragraph,
        text: str,
        base_bold: bool = False,
        base_italic: bool = False,
        parse_markdown: bool = True,
    ):
        """
        Додає текст до параграфа з підтримкою вбудованого Markdown:
        - **жирний**
        - *курсив*
        - ***жирний курсив***
        - `код` (Consolas + noProof для запобігання підкресленню орфографії)
        """
        if not text:
            return

        if not parse_markdown or not re.search(r'[`*$]', text):
            run = paragraph.add_run(text)
            self._apply_run_font(run, size_pt=self.font_size_body, bold=base_bold, italic=base_italic)
            return

        last_idx = 0
        for match in MD_INLINE_RE.finditer(text):
            start, end = match.span()
            if start > last_idx:
                chunk = text[last_idx:start]
                r = paragraph.add_run(chunk)
                self._apply_run_font(r, size_pt=self.font_size_body, bold=base_bold, italic=base_italic)

            group = match.lastgroup
            val = match.group(group)

            if group == "code":
                code_text = val[1:-1]
                r_code = paragraph.add_run(code_text)
                code_size = max(10.0, self.font_size_body - 1.0)
                self._apply_run_font(
                    r_code,
                    font_name="Consolas",
                    size_pt=code_size,
                    bold=False,
                    italic=False,
                    color_hex="000000",
                )
                rPr = r_code._r.get_or_add_rPr()
                rPr.append(parse_xml(f'<w:noProof {nsdecls("w")}/>'))
            elif group == "math":
                math_expr = val[1:-1]
                omml_xml = self.convert_to_omml(math_expr)
                if omml_xml:
                    paragraph._p.append(parse_xml(omml_xml))
                else:
                    r_math = paragraph.add_run(math_expr)
                    self._apply_run_font(r_math, size_pt=self.font_size_body, italic=True)
            elif group == "bold_italic":
                bi_text = val[3:-3]
                r_bi = paragraph.add_run(bi_text)
                self._apply_run_font(r_bi, size_pt=self.font_size_body, bold=True, italic=True)
            elif group == "bold":
                b_text = val[2:-2]
                r_b = paragraph.add_run(b_text)
                self._apply_run_font(r_b, size_pt=self.font_size_body, bold=True, italic=base_italic)
            elif group == "italic":
                i_text = val[1:-1]
                r_i = paragraph.add_run(i_text)
                self._apply_run_font(r_i, size_pt=self.font_size_body, bold=base_bold, italic=True)

            last_idx = end

        if last_idx < len(text):
            rem = text[last_idx:]
            r_rem = paragraph.add_run(rem)
            self._apply_run_font(r_rem, size_pt=self.font_size_body, bold=base_bold, italic=base_italic)

    # --------------------------------------------------------------------------
    # ПУБЛІЧНИЙ API ДЛЯ РОБОТИ З ТЕКСТОМ ТА СТРУКТУРОЮ
    # --------------------------------------------------------------------------

    def add_heading(
        self,
        text: str,
        level: int = 1,
        page_break_before: Optional[bool] = None,
        align: Optional[str] = None,
    ):
        """
        Додає заголовок H1-H3 із захистом від відриву від наступного тексту (keep_with_next).
        Вирівнювання береться з styles_config.json (з можливістю перевизначення параметром align).
        Для level=1 автоматично додається розрив сторінки (крім найпершого заголовка), якщо увімкнено h1_page_break_before.
        """
        should_break = page_break_before
        if should_break is None:
            if level == 1 and self.h1_page_break_before and self._h1_count > 0:
                should_break = True
            else:
                should_break = False

        if level == 1:
            self._h1_count += 1

        if should_break:
            self.add_page_break()

        sizes = {1: self.font_size_h1, 2: self.font_size_h2, 3: self.font_size_h3}
        size = sizes.get(level, self.font_size_h2)

        # Прив'язка до нативного стилю Word (Heading 1, Heading 2, Heading 3) для навігації та автозмісту
        style_name = f"Heading {level}" if f"Heading {level}" in self.doc.styles else "Heading 1"
        p = self.doc.add_paragraph(style=style_name)

        level_aligns = {
            1: self.heading_h1_align,
            2: self.heading_h2_align,
            3: self.heading_h3_align,
        }
        chosen_align = align or level_aligns.get(level, self.heading_align)
        p.alignment = ALIGN_MAP.get(str(chosen_align).lower(), WD_ALIGN_PARAGRAPH.LEFT)

        # Відступи заголовків
        p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
        p.paragraph_format.space_after = Pt(10 if level == 1 else 6)
        p.paragraph_format.keep_with_next = True

        run = p.add_run(text)
        self._apply_run_font(run, size_pt=size, bold=True)
        return p

    def add_paragraph(
        self,
        text: str,
        bold: bool = False,
        italic: bool = False,
        align: Optional[str] = None,
        indent: bool = True,
        space_after_pt: float = 0,
        parse_markdown: bool = True,
    ):
        """
        Додає абзац тексту з автоматичним абзацним відступом (ДСТУ) та вирівнюванням.
        Вирівнювання береться з styles_config.json (з можливістю перевизначення параметром align).
        Підтримує вбудований базовий Markdown (**жирний**, *курсив*, ***жирний курсив***, `код`).
        """
        p = self.doc.add_paragraph()
        align_map = {
            "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
        }
        chosen_align = align or self.paragraph_align
        p.alignment = align_map.get(str(chosen_align).lower(), WD_ALIGN_PARAGRAPH.JUSTIFY)
        p.paragraph_format.line_spacing = self.line_spacing
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after_pt)

        if indent:
            p.paragraph_format.first_line_indent = Cm(self.paragraph_indent_cm)

        self._add_formatted_text(
            p,
            text,
            base_bold=bold,
            base_italic=italic,
            parse_markdown=parse_markdown,
        )
        return p

    def add_formatted_paragraph(
        self,
        runs: List[Tuple[Any, ...]],
        align: Optional[str] = None,
        indent: bool = True,
        space_after_pt: float = 0,
    ):
        """
        Додає абзац, сформований зі списку кортежів для прямого контролю окремих Run:
        runs = [("текст", bold, italic), ("код", False, False, "Consolas", 12.0), ...]
        """
        p = self.doc.add_paragraph()
        align_map = {
            "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
        }
        chosen_align = align or self.paragraph_align
        p.alignment = align_map.get(str(chosen_align).lower(), WD_ALIGN_PARAGRAPH.JUSTIFY)
        p.paragraph_format.line_spacing = self.line_spacing
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after_pt)

        if indent:
            p.paragraph_format.first_line_indent = Cm(self.paragraph_indent_cm)

        for item in runs:
            text = str(item[0]) if len(item) > 0 else ""
            bold = bool(item[1]) if len(item) > 1 else False
            italic = bool(item[2]) if len(item) > 2 else False
            font_name = str(item[3]) if len(item) > 3 and item[3] else self.font_name
            size_pt = float(item[4]) if len(item) > 4 and item[4] else self.font_size_body
            color_hex = str(item[5]) if len(item) > 5 and item[5] else "000000"

            r = p.add_run(text)
            self._apply_run_font(
                r,
                font_name=font_name,
                size_pt=size_pt,
                bold=bold,
                italic=italic,
                color_hex=color_hex,
            )

        return p

    def add_bullet(
        self,
        text: str,
        bullet_char: str = "•",
        parse_markdown: bool = True,
        align: Optional[str] = None,
    ):
        """
        Додає пункт нативного маркованого списку Word (List Bullet).
        """
        # Створюємо параграф зі стандартним стилем маркованого списку Word
        p = self.doc.add_paragraph(style="List Bullet")
        chosen_align = align or self.bullet_align
        p.alignment = ALIGN_MAP.get(str(chosen_align).lower(), WD_ALIGN_PARAGRAPH.JUSTIFY)
        p.paragraph_format.line_spacing = self.line_spacing
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        # Налаштовуємо відступи (відступ тексту 1.25 см)
        p.paragraph_format.left_indent = Cm(self.paragraph_indent_cm)
        # Додаємо лише сам текст (без вставки символу "• ", оскільки Word додасть його автоматично!)
        self._add_formatted_text(p, text, parse_markdown=parse_markdown)
        return p

    def add_numbered(
        self,
        text: str,
        parse_markdown: bool = True,
        align: Optional[str] = None,
    ):
        """
        Додає пункт нативного нумерованого списку Word (List Number).
        """
        p = self.doc.add_paragraph(style="List Number")
        chosen_align = align or self.numbered_align
        p.alignment = ALIGN_MAP.get(str(chosen_align).lower(), WD_ALIGN_PARAGRAPH.JUSTIFY)
        p.paragraph_format.line_spacing = self.line_spacing
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Cm(self.paragraph_indent_cm)
        self._add_formatted_text(p, text, parse_markdown=parse_markdown)
        return p

    def _create_caption_paragraph(
        self,
        caption: str,
        position: str = "below",
        align: str = "center",
        italic: bool = True,
    ):
        """
        Створює параграф підпису для таблиці чи зображення із врахуванням позиції та вирівнювання.
        Якщо позиція 'above' — автоматично вмикає keep_with_next для захисту від відриву від таблиці/зображення.
        """
        p_cap = self.doc.add_paragraph()
        align_map = {
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
            "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        }
        p_cap.alignment = align_map.get(align.lower(), WD_ALIGN_PARAGRAPH.CENTER)
        if position.lower() == "above":
            p_cap.paragraph_format.space_before = Pt(6)
            p_cap.paragraph_format.space_after = Pt(3)
            p_cap.paragraph_format.keep_with_next = True
        else:
            p_cap.paragraph_format.space_before = Pt(3)
            p_cap.paragraph_format.space_after = Pt(6)

        self._add_formatted_text(p_cap, caption, base_italic=italic, parse_markdown=True)
        return p_cap

    def add_image(
        self,
        image_name_or_path: str,
        caption: Optional[str] = None,
        width_cm: Optional[float] = None,
        caption_position: Optional[str] = None,
        caption_align: Optional[str] = None,
    ):
        """
        Додає фото із папки images/ або прямого шляху:
        - Автоматично масштабує під ширину полів сторінки.
        - Розташування підпису ("above" або "below") береться з styles_config.json або параметра caption_position.
        """
        target_path = image_name_or_path
        if not os.path.exists(target_path):
            candidate = os.path.join(self.images_dir, image_name_or_path)
            if os.path.exists(candidate):
                target_path = candidate
            else:
                current_dir = os.path.dirname(os.path.abspath(__file__))
                candidate2 = os.path.join(current_dir, self.images_dir, image_name_or_path)
                if os.path.exists(candidate2):
                    target_path = candidate2

        if not os.path.exists(target_path):
            p_err = self.doc.add_paragraph()
            p_err.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_err = p_err.add_run(f"[Зображення не знайдено: {image_name_or_path}]")
            self._apply_run_font(r_err, size_pt=12.0, italic=True, color_hex="FF0000")
            return None

        # Розрахунок максимальної доступної ширини (A4 21.0 см - поля)
        max_width_cm = 21.0 - self.margin_left_cm - self.margin_right_cm
        final_width = Cm(width_cm) if (width_cm and width_cm <= max_width_cm) else Cm(max_width_cm)

        pos = (caption_position or self.image_caption_position).lower()
        aln = (caption_align or self.image_caption_align).lower()

        # Підпис зверху (якщо pos == "above")
        if caption and pos == "above":
            self._create_caption_paragraph(caption, position="above", align=aln, italic=True)

        # Вставка картинки
        p_img = self.doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(4 if pos == "above" else 6)
        p_img.paragraph_format.space_after = Pt(6 if pos == "above" else 2)
        if pos == "below":
            p_img.paragraph_format.keep_with_next = True

        run_img = p_img.add_run()
        run_img.add_picture(target_path, width=final_width)

        # Підпис знизу (якщо pos == "below")
        if caption and pos == "below":
            self._create_caption_paragraph(caption, position="below", align=aln, italic=True)

        return p_img

    def add_table(
        self,
        headers: List[str],
        rows: List[List[str]],
        caption: Optional[str] = None,
        col_widths_cm: Optional[List[float]] = None,
        caption_position: Optional[str] = None,
        caption_align: Optional[str] = None,
    ):
        """
        Додає академічну/технічну таблицю:
        - Суцільні одинарні межі навколо кожної комірки (1 pt).
        - Внутрішні відступи комірок (padding).
        - Захист рядків від розриву навпіл між сторінками (cantSplit).
        - Автоматичне дублювання шапки на нових сторінках (tblHeader).
        - Розташування підпису ("above" або "below") береться з styles_config.json або параметра caption_position.
        """
        pos = (caption_position or self.table_caption_position).lower()
        aln = (caption_align or self.table_caption_align).lower()

        # Підпис зверху (якщо pos == "above")
        if caption and pos == "above":
            self._create_caption_paragraph(caption, position="above", align=aln, italic=False)

        num_rows = len(rows) + 1
        num_cols = len(headers)

        table = self.doc.add_table(rows=num_rows, cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        # Розрахунок ширини колонок
        max_table_width = 21.0 - self.margin_left_cm - self.margin_right_cm
        if col_widths_cm and len(col_widths_cm) == num_cols:
            col_widths = [Cm(w) for w in col_widths_cm]
        else:
            col_widths = [Cm(max_table_width / num_cols)] * num_cols

        # Налаштування шапки (перший рядок)
        header_row = table.rows[0]
        self._set_repeat_header(header_row)
        self._set_cant_split(header_row)

        for c_idx, text in enumerate(headers):
            cell = header_row.cells[c_idx]
            cell.width = col_widths[c_idx]
            self._set_cell_borders(cell, border_color=self.table_border_color, sz=self.table_border_sz)
            self._set_cell_margins(cell, top_twips=100, bottom_twips=100, left_twips=140, right_twips=140)

            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            self._add_formatted_text(p, str(text), base_bold=True, parse_markdown=True)

        # Рядки даних
        for r_idx, row_data in enumerate(rows):
            row = table.rows[r_idx + 1]
            self._set_cant_split(row)

            for c_idx, val in enumerate(row_data):
                cell = row.cells[c_idx]
                cell.width = col_widths[c_idx]
                self._set_cell_borders(cell, border_color=self.table_border_color, sz=self.table_border_sz)
                self._set_cell_margins(cell, top_twips=100, bottom_twips=100, left_twips=140, right_twips=140)

                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                self._add_formatted_text(p, str(val), parse_markdown=True)

        # Підпис знизу (якщо pos == "below")
        if caption and pos == "below":
            self._create_caption_paragraph(caption, position="below", align=aln, italic=False)

        return table

    def add_toc(self, title: str = "ЗМІСТ"):
        """Вставляє динамічне поле автоматичного змісту (Table of Contents)."""
        if title:
            # Заголовок блоку змісту без outline level (щоб сам "ЗМІСТ" не дублювався у змісті)
            p_title = self.doc.add_paragraph()
            p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_title.paragraph_format.space_before = Pt(0)
            p_title.paragraph_format.space_after = Pt(12)
            p_title.paragraph_format.keep_with_next = True
            r_title = p_title.add_run(title)
            self._apply_run_font(r_title, size_pt=self.font_size_h1, bold=True)

        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(12)
        r = p.add_run()
        self._insert_field_code(r, r'TOC \o "1-3" \h \z \u')
        return p

    def add_code_block(self, code: str, caption: Optional[str] = None):
        """
        Додає блок коду зі збереженням структури рядків та вимкненням перевірки орфографії (noProof).
        """
        lines = code.split("\n")
        for line in lines:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0

            run = p.add_run(line if line.strip() else " ")
            self._apply_run_font(run, font_name="Consolas", size_pt=11.0)

            rPr = run._r.get_or_add_rPr()
            rPr.append(parse_xml(f'<w:noProof {nsdecls("w")}/>'))

        if caption:
            p_cap = self.doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(6)
            run_cap = p_cap.add_run(caption)
            self._apply_run_font(run_cap, size_pt=self.font_size_body, italic=True)

    def add_formula(
        self,
        formula: str,
        number: Optional[str] = None,
        space_before_pt: float = 6,
        space_after_pt: float = 8,
    ):
        """
        Додає формулу у форматі LaTeX або MathML:
        - Конвертує у нативний OpenXML oMath об'єкт Word (редаговані дроби, корені, суми, інтеграли).
        - Якщо вказано number (наприклад '1.1') — вирівнює формулу по центру, а номер праворуч у дужках (1.1).
        - Має багаторівневий fallback: OMML -> Matplotlib PNG -> Текстовий вираз.
        """
        max_w = 21.0 - self.margin_left_cm - self.margin_right_cm
        omml_xml = self.convert_to_omml(formula)

        if omml_xml:
            if number:
                table = self.doc.add_table(rows=1, cols=2)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = False

                c0 = table.rows[0].cells[0]
                c1 = table.rows[0].cells[1]
                c0.width = Cm(max_w - 2.5)
                c1.width = Cm(2.5)

                p0 = c0.paragraphs[0]
                p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p0.paragraph_format.space_before = Pt(space_before_pt)
                p0.paragraph_format.space_after = Pt(space_after_pt)
                p0._p.append(parse_xml(omml_xml))

                p1 = c1.paragraphs[0]
                p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                p1.paragraph_format.space_before = Pt(space_before_pt)
                p1.paragraph_format.space_after = Pt(space_after_pt)
                r1 = p1.add_run(f"({number})")
                self._apply_run_font(r1, size_pt=self.font_size_body)
                return table
            else:
                p = self.doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(space_before_pt)
                p.paragraph_format.space_after = Pt(space_after_pt)
                p._p.append(parse_xml(omml_xml))
                return p

        # Fallback 1: Matplotlib PNG
        png_path = self._render_latex_to_temp_png(formula)
        if png_path and os.path.exists(png_path):
            try:
                if number:
                    table = self.doc.add_table(rows=1, cols=2)
                    table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    table.autofit = False

                    c0 = table.rows[0].cells[0]
                    c1 = table.rows[0].cells[1]
                    c0.width = Cm(max_w - 2.5)
                    c1.width = Cm(2.5)

                    p0 = c0.paragraphs[0]
                    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p0.paragraph_format.space_before = Pt(space_before_pt)
                    p0.paragraph_format.space_after = Pt(space_after_pt)
                    r0 = p0.add_run()
                    r0.add_picture(png_path)

                    p1 = c1.paragraphs[0]
                    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    p1.paragraph_format.space_before = Pt(space_before_pt)
                    p1.paragraph_format.space_after = Pt(space_after_pt)
                    r1 = p1.add_run(f"({number})")
                    self._apply_run_font(r1, size_pt=self.font_size_body)
                    return table
                else:
                    p = self.doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(space_before_pt)
                    p.paragraph_format.space_after = Pt(space_after_pt)
                    r = p.add_run()
                    r.add_picture(png_path)
                    return p
            finally:
                if os.path.exists(png_path):
                    try:
                        os.remove(png_path)
                    except Exception:
                        pass

        # Fallback 2: Звичайний текст
        if number:
            table = self.doc.add_table(rows=1, cols=2)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False

            c0 = table.rows[0].cells[0]
            c1 = table.rows[0].cells[1]
            c0.width = Cm(max_w - 2.5)
            c1.width = Cm(2.5)

            p0 = c0.paragraphs[0]
            p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p0.paragraph_format.space_before = Pt(space_before_pt)
            p0.paragraph_format.space_after = Pt(space_after_pt)
            r0 = p0.add_run(formula)
            self._apply_run_font(r0, size_pt=self.font_size_h2, bold=True)

            p1 = c1.paragraphs[0]
            p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p1.paragraph_format.space_before = Pt(space_before_pt)
            p1.paragraph_format.space_after = Pt(space_after_pt)
            r1 = p1.add_run(f"({number})")
            self._apply_run_font(r1, size_pt=self.font_size_body)
            return table
        else:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(space_before_pt)
            p.paragraph_format.space_after = Pt(space_after_pt)
            run = p.add_run(formula)
            self._apply_run_font(run, size_pt=self.font_size_h2, bold=True)
            return p

    def add_page_break(self):
        """Додає розрив сторінки."""
        self.doc.add_page_break()

    def save(self, file_path: str) -> str:
        """
        Зберігає документ за вказаним шляхом.
        Якщо файл заблоковано іншою програмою (наприклад, відкрито у Microsoft Word),
        перехоплює PermissionError / Errno 13 та автоматично зберігає у резервний файл
        (<шлях>_new.docx або <шлях>_new_1.docx тощо), виводячи детальне та зрозуміле попередження.
        Повертає фактичний шлях до збереженого файлу.
        """
        try:
            self.doc.save(file_path)
            print(f"✅ Документ успішно створено: {file_path}")
            return file_path
        except (PermissionError, OSError) as e:
            # Перевіряємо блокування доступу у Windows (Errno 13 Permission denied)
            if isinstance(e, PermissionError) or getattr(e, "errno", None) == 13:
                filename = os.path.basename(file_path)
                print(f"\n⚠️  ПОПЕРЕДЖЕННЯ: Файл '{filename}' заблоковано системою Windows.")
                print("   Найімовірніше, цей документ зараз відкрито у Microsoft Word.")

                base, ext = os.path.splitext(file_path)
                fallback_path = f"{base}_new{ext}"
                counter = 1
                while counter <= 100:
                    try:
                        self.doc.save(fallback_path)
                        print(f"📁 Документ успішно збережено в резервний файл:\n   -> {fallback_path}")
                        print(f"💡 Щоб перезаписати основний файл, закрийте '{filename}' у Word перед наступною збіркою.\n")
                        return fallback_path
                    except (PermissionError, OSError):
                        fallback_path = f"{base}_new_{counter}{ext}"
                        counter += 1

                # Аварійний резерв у системну тимчасову папку
                temp_dir = tempfile.gettempdir()
                temp_fallback = os.path.join(temp_dir, f"{os.path.splitext(filename)[0]}_{uuid.uuid4().hex[:6]}{ext}")
                self.doc.save(temp_fallback)
                print(f"📁 Документ аварійно збережено у тимчасову папку:\n   -> {temp_fallback}\n")
                return temp_fallback
            raise e

    # --------------------------------------------------------------------------
    # ВНУТРІШНІ ДОПОМІЖНІ МЕТОДИ OPENXML
    # --------------------------------------------------------------------------

    def _set_cell_borders(self, cell, border_color: str = "000000", sz: str = "8"):
        tcPr = cell._tc.get_or_add_tcPr()
        c = clean_hex(border_color)
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{c}"/>'
            f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{c}"/>'
            f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{c}"/>'
            f'<w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{c}"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)

    def _set_cell_margins(
        self,
        cell,
        top_twips=100,
        bottom_twips=100,
        left_twips=140,
        right_twips=140,
    ):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="{top_twips}" w:type="dxa"/>'
            f'<w:bottom w:w="{bottom_twips}" w:type="dxa"/>'
            f'<w:left w:w="{left_twips}" w:type="dxa"/>'
            f'<w:right w:w="{right_twips}" w:type="dxa"/>'
            f'</w:tcMar>'
        )
        tcPr.append(tcMar)

    def _set_repeat_header(self, row):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

    def _set_cant_split(self, row):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
